import io
from datetime import datetime, timezone

from openpyxl import load_workbook

from app.models import Assessment, AssessmentStatus, Result, UserRole
from tests.conftest import auth_token, create_class, create_user


async def _complete_via_http(client, headers, scale_id: str) -> int:
    """跑完一遍 MHT（0/1 计分）并返回 assessment_id。"""
    start = await client.post("/api/assess/start", json={"scale_id": scale_id}, headers=headers)
    assert start.status_code == 200, start.text
    aid = start.json()["assessment_id"]
    for _ in range(200):
        nxt = await client.get(f"/api/assess/{aid}/next", headers=headers)
        data = nxt.json()
        if data["finished"]:
            break
        resp = await client.post(
            f"/api/assess/{aid}/answer",
            json={"item_id": data["item_id"], "value": 0},
            headers=headers,
        )
        assert resp.status_code == 200, resp.text
    finish = await client.post(f"/api/assess/{aid}/finish", headers=headers)
    assert finish.status_code == 200, finish.text
    return aid


async def _make_completed(db, student, scale_id: str, scores: dict, *, is_valid=True,
                          risk_flags=None, attempt: int = 0, llm_summary=None,
                          fallback_used=True) -> Assessment:
    when = datetime(2026, 10, 15, 10, attempt, tzinfo=timezone.utc)
    assessment = Assessment(
        user_id=student.id,
        scale_id=scale_id,
        status=AssessmentStatus.COMPLETED,
        started_at=when,
        completed_at=when,
        is_valid=is_valid,
    )
    db.add(assessment)
    await db.flush()
    db.add(
        Result(
            assessment_id=assessment.id,
            scores=scores,
            risk_flags=risk_flags or [],
            llm_summary=llm_summary,
            fallback_used=fallback_used,
            created_at=when,
        )
    )
    await db.commit()
    await db.refresh(assessment)
    return assessment


# ---------------------------------------------------------------- 问题 1 / 2

async def test_my_results_returns_completed_for_mht_student(client, db):
    cls = await create_class(db, "九年级(7)班")
    student = await create_user(db, UserRole.STUDENT, "stu_report", class_id=cls.id)
    token = await auth_token(client, "stu_report")
    headers = {"Authorization": f"Bearer {token}"}

    await _complete_via_http(client, headers, "mht")

    resp = await client.get("/api/assess/my-results", headers=headers)
    assert resp.status_code == 200
    results = resp.json()["results"]
    assert len(results) >= 1
    completed = [r for r in results if r["status"] == "completed"]
    assert completed, "至少应有一条 completed"
    # 中文量表名，而非裸 scale_id
    assert completed[0]["scale_name"] != "mht"
    assert completed[0]["attempt_no"] == 1
    assert "scores" in completed[0]


async def test_scales_attempt_fields_and_multi_attempt(client, db):
    cls = await create_class(db, "九年级(8)班")
    student = await create_user(db, UserRole.STUDENT, "stu_multi", class_id=cls.id)
    token = await auth_token(client, "stu_multi")
    headers = {"Authorization": f"Bearer {token}"}

    first = await _complete_via_http(client, headers, "mht")

    scales = await client.get("/api/assess/scales", headers=headers)
    mht = next(s for s in scales.json() if s["scale_id"] == "mht")
    assert mht["completed"] is True
    assert mht["attempts"] == 1
    assert mht["in_progress"] is False
    assert mht["completed_at"] is not None

    # 开展第二次测评（无进行中的记录 → 新建一次尝试）
    start = await client.post("/api/assess/start", json={"scale_id": "mht"}, headers=headers)
    assert start.status_code == 200
    second = start.json()["assessment_id"]
    assert second != first

    # 尚未完成时 in_progress=True
    scales = await client.get("/api/assess/scales", headers=headers)
    mht = next(s for s in scales.json() if s["scale_id"] == "mht")
    assert mht["in_progress"] is True
    assert mht["attempts"] == 1

    # 完成第二次
    for _ in range(200):
        nxt = await client.get(f"/api/assess/{second}/next", headers=headers)
        data = nxt.json()
        if data["finished"]:
            break
        await client.post(
            f"/api/assess/{second}/answer",
            json={"item_id": data["item_id"], "value": 0},
            headers=headers,
        )
    finish = await client.post(f"/api/assess/{second}/finish", headers=headers)
    assert finish.status_code == 200

    scales = await client.get("/api/assess/scales", headers=headers)
    mht = next(s for s in scales.json() if s["scale_id"] == "mht")
    assert mht["attempts"] == 2
    assert mht["completed"] is True
    assert mht["in_progress"] is False

    results = (await client.get("/api/assess/my-results", headers=headers)).json()["results"]
    mht_results = [r for r in results if r["scale_id"] == "mht"]
    assert len(mht_results) == 2
    assert {r["attempt_no"] for r in mht_results} == {1, 2}


# ---------------------------------------------------------------- 问题 3

async def test_classes_report_has_pending_counts(client, db):
    cls = await create_class(db, "九年级(7)班")
    s1 = await create_user(db, UserRole.STUDENT, "stu_cls1", class_id=cls.id)
    await create_user(db, UserRole.STUDENT, "stu_cls2", class_id=cls.id)
    await _make_completed(
        db, s1, "mht",
        {"total_score": 30, "total_label": "正常", "validity_score": 1, "is_valid": True,
         "subscales": {"phobia": {"name": "恐怖倾向", "score": 3}}},
    )
    admin = await create_user(db, UserRole.SCHOOL_ADMIN, "admin_cls")
    token = await auth_token(client, "admin_cls")

    resp = await client.get("/api/admin/reports/classes", headers={"Authorization": f"Bearer {token}"})
    assert resp.status_code == 200
    row = next(r for r in resp.json() if r["class_id"] == cls.id)
    assert row["total"] == 2
    assert row["completed"] == 1
    assert row["pending"] == 1
    assert row["completion_rate"] == 0.5


async def test_admin_export_includes_pending_and_no_teacher_data(client, db):
    cls = await create_class(db, "九年级(7)班")
    done = await create_user(db, UserRole.STUDENT, "stu_done", real_name="完成同学", class_id=cls.id)
    await create_user(db, UserRole.STUDENT, "stu_pending", real_name="未完成同学", class_id=cls.id)
    await _make_completed(
        db, done, "mht",
        {"total_score": 30, "total_label": "正常", "validity_score": 1, "is_valid": True,
         "subscales": {}},
        risk_flags=[{"severity": "high", "message": "x"}],
    )
    # 同班教师 + 教师量表结果：导出绝不能包含
    teacher = await create_user(db, UserRole.TEACHER, "tea_exp", real_name="教师甲",
                                class_id=cls.id, is_homeroom=True)
    await _make_completed(
        db, teacher, "teacher_ref",
        {"total_mean": 4.0, "factors": {"self": {"name": "自我", "mean": 4.0}}},
    )
    admin = await create_user(db, UserRole.SCHOOL_ADMIN, "admin_exp")
    token = await auth_token(client, "admin_exp")
    headers = {"Authorization": f"Bearer {token}"}

    resp = await client.get(
        f"/api/admin/reports/export?class_id={cls.id}&scale_id=mht&format=xlsx",
        headers=headers,
    )
    assert resp.status_code == 200, resp.text
    assert "attachment" in resp.headers["content-disposition"]

    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header, data = rows[0], rows[1:]
    assert len(data) == 2, "行数应等于该班学生数（含未完成）"
    names = [r[0] for r in data]
    assert "教师甲" not in names, "管理员导出不得包含教师数据"
    assert "完成同学" in names and "未完成同学" in names
    # 新列顺序：姓名|班级|性别|状态|量表|第几次|有效性(效度)|...
    assert header[:7] == ("姓名", "班级", "性别", "状态", "量表", "第几次", "有效性(效度)")
    status_by_name = {r[0]: r[3] for r in data}
    assert status_by_name["未完成同学"] == "未完成"
    assert status_by_name["完成同学"] == "已完成"
    assert "状态" in header and "有效性(效度)" in header
    assert "LLM 测评结论" in header


async def test_export_includes_llm_conclusion_column(client, db):
    """问题 3：导出新增「LLM 测评结论」列；fallback 追加「（模板生成）」；未完成留空。"""
    cls = await create_class(db, "九年级(9)班")
    with_llm = await create_user(db, UserRole.STUDENT, "stu_llm", real_name="有结论同学", class_id=cls.id)
    fallback = await create_user(db, UserRole.STUDENT, "stu_llm_fb", real_name="模板同学", class_id=cls.id)
    await create_user(db, UserRole.STUDENT, "stu_llm_pending", real_name="待完成同学", class_id=cls.id)
    base_scores = {"total_score": 12, "total_label": "正常", "validity_score": 1, "is_valid": True,
                   "subscales": {}}
    await _make_completed(
        db, with_llm, "mht", base_scores,
        llm_summary="该生整体状态平稳，建议保持规律作息。", fallback_used=False,
    )
    await _make_completed(
        db, fallback, "mht", base_scores,
        llm_summary="模板结论：整体平稳。", fallback_used=True,
    )
    admin = await create_user(db, UserRole.SCHOOL_ADMIN, "admin_llm_col")
    token = await auth_token(client, "admin_llm_col")
    headers = {"Authorization": f"Bearer {token}"}

    resp = await client.get(
        f"/api/admin/reports/export?class_id={cls.id}&scale_id=mht&format=xlsx",
        headers=headers,
    )
    assert resp.status_code == 200, resp.text
    wb = load_workbook(io.BytesIO(resp.content))
    header = list(next(wb.active.iter_rows(values_only=True)))
    idx = header.index("LLM 测评结论")
    # 统一列顺序：LLM 结论在「预警级别」之后、「完成时间」之前。
    assert header[-3:] == ["预警级别", "LLM 测评结论", "完成时间"]

    data = list(wb.active.iter_rows(values_only=True))[1:]
    by_name = {r[0]: r for r in data}
    assert by_name["有结论同学"][idx] == "该生整体状态平稳，建议保持规律作息。"
    assert by_name["模板同学"][idx] == "模板结论：整体平稳。（模板生成）"
    assert not by_name["待完成同学"][idx], "未完成学生的 LLM 结论应为空"


async def test_teacher_export_scope_and_forbidden(client, db):
    cls_a = await create_class(db, "九年级(1)班")
    cls_b = await create_class(db, "九年级(2)班")
    teacher = await create_user(db, UserRole.TEACHER, "tea_a", real_name="甲老师",
                                class_id=cls_a.id, is_homeroom=True)
    cls_a.head_teacher_id = teacher.id
    await create_user(db, UserRole.TEACHER, "tea_b", real_name="乙老师",
                      class_id=cls_b.id, is_homeroom=True)
    student_a = await create_user(db, UserRole.STUDENT, "stu_a1", real_name="甲班学生", class_id=cls_a.id)
    await create_user(db, UserRole.STUDENT, "stu_a2", real_name="甲班待完成", class_id=cls_a.id)
    await _make_completed(
        db, student_a, "mht",
        {"total_score": 40, "total_label": "欠佳", "validity_score": 1, "is_valid": True,
         "subscales": {}},
    )
    await db.commit()

    token = await auth_token(client, "tea_a")
    headers = {"Authorization": f"Bearer {token}"}

    summary = await client.get("/api/teacher/class-summary?scale_id=mht", headers=headers)
    assert summary.status_code == 200, summary.text
    body = summary.json()
    assert body["class_name"] == "九年级(1)班"
    assert body["total"] == 2
    assert body["completed"] == 1
    assert body["pending"] == 1

    own = await client.get("/api/teacher/class-export?scale_id=mht&format=xlsx", headers=headers)
    assert own.status_code == 200, own.text
    wb = load_workbook(io.BytesIO(own.content))
    data = list(wb.active.iter_rows(values_only=True))[1:]
    assert len(data) == 2
    assert {r[0] for r in data} == {"甲班学生", "甲班待完成"}

    # 他班（显式 class_id）→ 403
    other = await client.get(
        f"/api/teacher/class-export?scale_id=mht&format=xlsx&class_id={cls_b.id}",
        headers=headers,
    )
    assert other.status_code == 403

    # 教师不能调用管理员导出
    forbidden = await client.get(
        f"/api/admin/reports/export?class_id={cls_a.id}&scale_id=mht&format=xlsx",
        headers=headers,
    )
    assert forbidden.status_code == 403


# ---------------------------------------------------------------- 任务书 12：统计口径与导出对齐

async def _setup_scale_alignment_class(db, name: str = "九年级(8)班"):
    """复现任务书实证：一名学生 MSSMHS 已完成 / MHT 进行中，另两名待完成。"""
    cls = await create_class(db, name)
    a = await create_user(db, UserRole.STUDENT, "s24jifenglan", real_name="纪凤兰", class_id=cls.id)
    await create_user(db, UserRole.STUDENT, "stu12_b", real_name="乙同学", class_id=cls.id)
    await create_user(db, UserRole.STUDENT, "stu12_c", real_name="丙同学", class_id=cls.id)
    await _make_completed(
        db, a, "mssmhs",
        {"total_mean": 2.0, "overall_level": "normal",
         "factors": {"self": {"name": "自我", "mean": 2.0}}},
        llm_summary="该生整体平稳。", fallback_used=False,
    )
    db.add(Assessment(user_id=a.id, scale_id="mht", status=AssessmentStatus.IN_PROGRESS))
    await db.commit()
    return cls


async def test_classes_report_scale_id_alignment(client, db):
    """任务书 12 第三节 1/2：mht/mssmhs/all 分别统计；非法值 400。"""
    cls = await _setup_scale_alignment_class(db)
    admin = await create_user(db, UserRole.SCHOOL_ADMIN, "admin12a")
    token = await auth_token(client, "admin12a")
    headers = {"Authorization": f"Bearer {token}"}

    async def row_for(scale: str) -> dict:
        resp = await client.get(
            f"/api/admin/reports/classes?scale_id={scale}", headers=headers
        )
        assert resp.status_code == 200, resp.text
        return next(r for r in resp.json() if r["class_id"] == cls.id)

    mht = await row_for("mht")
    mssmhs = await row_for("mssmhs")
    allr = await row_for("all")
    print(
        f"[verify:classes] class_id={cls.id} completed -> "
        f"mht={mht['completed_count']} mssmhs={mssmhs['completed_count']} "
        f"all={allr['completed_count']} | pending -> mht={mht['pending']} "
        f"mssmhs={mssmhs['pending']} all={allr['pending']}"
    )
    # 任务书第三节第 1 条：实测应为 mht=0、mssmhs=1、all=1。
    assert mht["completed_count"] == 0
    assert mssmhs["completed_count"] == 1
    assert allr["completed_count"] == 1
    assert (mht["pending"], mssmhs["pending"], allr["pending"]) == (3, 2, 2)
    # dimension_means 也按选中量表：mssmhs 有 self 因子，mht 无已完成 → 空。
    assert mssmhs["dimension_means"].get("self") == 2.0
    assert mht["dimension_means"] == {}

    bogus = await client.get("/api/admin/reports/classes?scale_id=bogus", headers=headers)
    print(f"[verify:classes] scale_id=bogus -> {bogus.status_code} {bogus.json()}")
    assert bogus.status_code == 400


async def test_export_scale_matches_class_completion(client, db):
    """任务书 12 第三节 3：导出 mht 0 行完成；mssmhs 1 行完成且含 LLM 结论列内容。"""
    cls = await _setup_scale_alignment_class(db, "九年级(8)班")
    admin = await create_user(db, UserRole.SCHOOL_ADMIN, "admin12b")
    token = await auth_token(client, "admin12b")
    headers = {"Authorization": f"Bearer {token}"}

    for scale, expected_completed in (("mht", 0), ("mssmhs", 1)):
        resp = await client.get(
            f"/api/admin/reports/export?class_id={cls.id}&scale_id={scale}&format=xlsx",
            headers=headers,
        )
        assert resp.status_code == 200, resp.text
        wb = load_workbook(io.BytesIO(resp.content))
        rows = list(wb.active.iter_rows(values_only=True))
        header, data = rows[0], rows[1:]
        status_idx = header.index("状态")
        llm_idx = header.index("LLM 测评结论")
        completed = [r for r in data if r[status_idx] == "已完成"]
        print(
            f"[verify:export] class_id={cls.id} scale_id={scale} "
            f"rows={len(data)} completed_rows={len(completed)}"
        )
        assert len(completed) == expected_completed, "导出口径须与班级统计一致"
        if expected_completed == 1:
            assert completed[0][0] == "纪凤兰"
            assert completed[0][llm_idx] == "该生整体平稳。"


# ---------------------------------------------------------------- 问题 1 / 2 自动可见

async def test_teacher_class_results_auto_visible_with_student_info(client, db):
    """学生答完即自动可见（无需分发），且返回学生姓名/用户名/班级/性别/第几次/中文量表名。"""
    cls = await create_class(db, "九年级(3)班")
    teacher = await create_user(db, UserRole.TEACHER, "tea_info", real_name="信息老师",
                                class_id=cls.id, is_homeroom=True)
    cls.head_teacher_id = teacher.id
    student = await create_user(db, UserRole.STUDENT, "stu_info", real_name="张三",
                                class_id=cls.id)
    await _make_completed(
        db, student, "mht",
        {"total_score": 20, "total_label": "正常", "validity_score": 1, "is_valid": True,
         "subscales": {}},
        llm_summary="结论文字", fallback_used=False,
    )
    # 同班教师 + 教师量表结果：绝不能出现在本班学生结果里。
    await _make_completed(
        db, teacher, "teacher_ref",
        {"total_mean": 4.0, "factors": {"self": {"name": "自我", "mean": 4.0}}},
    )
    await db.commit()

    token = await auth_token(client, "tea_info")
    headers = {"Authorization": f"Bearer {token}"}

    resp = await client.get("/api/teacher/class-results?scale_id=mht", headers=headers)
    assert resp.status_code == 200, resp.text
    rows = resp.json()
    assert len(rows) == 1, "只应有本班学生（教师量表结果不得出现）"
    row = rows[0]
    assert row["student_name"] == "张三"
    assert row["username"] == "stu_info"
    assert row["class_name"] == cls.name
    assert row["gender"] == "男"
    assert row["scale_id"] == "mht"
    assert row["scale_name"] != "mht", "必须是中文量表名"
    assert row["attempt_no"] == 1
    assert row["completed_at"] is not None
    assert row["llm_summary"] == "结论文字"

    # 任务书点名的路径同样可用（自动可见，不再依赖分发）。
    alias = await client.get("/api/teacher/class-students-results?scale_id=mht", headers=headers)
    assert alias.status_code == 200, alias.text
    assert len(alias.json()) == 1


async def test_teacher_class_results_forbidden_for_non_homeroom(client, db):
    cls = await create_class(db, "九年级(4)班")
    homeroom = await create_user(db, UserRole.TEACHER, "tea_homeroom", real_name="班主",
                                 class_id=cls.id, is_homeroom=True)
    cls.head_teacher_id = homeroom.id
    other = await create_user(db, UserRole.TEACHER, "tea_other", real_name="非班主",
                              class_id=cls.id, is_homeroom=False)
    await db.commit()

    token = await auth_token(client, "tea_other")
    resp = await client.get(
        "/api/teacher/class-results?scale_id=mht",
        headers={"Authorization": f"Bearer {token}"},
    )
    assert resp.status_code == 403

async def test_my_results_includes_own_student_info(client, db):
    """报告页要展示学生本人信息：my-results 必须带 student，且 /me 与自己一致（不泄露他人）。"""
    cls = await create_class(db, "九年级(9)班")
    student = await create_user(db, UserRole.STUDENT, "stu_brief", class_id=cls.id)
    token = await auth_token(client, "stu_brief")
    headers = {"Authorization": f"Bearer {token}"}
    await _complete_via_http(client, headers, "mht")

    data = (await client.get("/api/assess/my-results", headers=headers)).json()
    assert data.get("student"), "缺少 student 信息块"
    st = data["student"]
    for key in ("id", "real_name", "username", "gender", "class_name"):
        assert key in st, f"student 缺少字段 {key}"
    assert st["username"] == "stu_brief"
    assert st["class_name"] == "九年级(9)班"
    assert st["id"] == student.id

    me = (await client.get("/api/auth/me", headers=headers)).json()
    assert (me["id"], me["username"], me["class_name"]) == (st["id"], st["username"], st["class_name"])
    assert me["role"] == "student"
