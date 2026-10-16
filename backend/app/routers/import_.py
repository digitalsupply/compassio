import csv
import io
import random
import re
import string
from datetime import datetime
from typing import Any

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile, status
from openpyxl import load_workbook
from pypinyin import lazy_pinyin
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.db import get_db
from app.models import Class, GradeLevel, User, UserRole
from app.observability.metrics import import_created_total
from app.schemas import ImportCommitResult, ImportPreview, ImportRowError
from app.security import hash_password, require_admin

router = APIRouter()

_charset = "abcdefghjkmnpqrstuvwxyzABCDEFGHJKMNPQRSTUVWXYZ23456789!@#$%&*"


def _gen_password(length: int = 8) -> str:
    return "".join(random.choice(_charset) for _ in range(length))


def _pinyin(name: str) -> str:
    return "".join(lazy_pinyin(name)).lower()


_JUNIOR_ALIAS = {"初一": "七年级", "初二": "八年级", "初三": "九年级"}
_CN_PRIMARY = "一二三四五六"


def _normalize_class_name(raw: str) -> str:
    """把班级名统一成 '七年级(1)班' 这样的规范格式。

    兼容常见的旧写法：初一(3)班 → 七年级(3)班；全角括号变半角。
    """
    name = (raw or "").strip().replace("（", "(").replace("）", ")")
    for old, new in _JUNIOR_ALIAS.items():
        if name.startswith(old):
            return new + name[len(old) :]
    return name


def _infer_grade_level(class_name: str) -> tuple[GradeLevel, bool]:
    name = _normalize_class_name(class_name)
    # 先看明确的学段关键字
    if "小" in name:
        return GradeLevel.PRIMARY_G4, False
    if "初" in name:
        return GradeLevel.JUNIOR, False
    if "高" in name:
        return GradeLevel.SENIOR, False
    # 再按年级数字/汉字推断
    if re.search(rf"[{_CN_PRIMARY}]年级", name) or re.match(r"^[1-6]年级", name):
        return GradeLevel.PRIMARY_G4, False
    if re.search(r"七|八|九", name):
        return GradeLevel.JUNIOR, False
    return GradeLevel.JUNIOR, True


async def _resolve_class(db: AsyncSession, class_name: str) -> Class:
    class_name = _normalize_class_name(class_name)
    result = await db.execute(select(Class).where(Class.name == class_name))
    cls = result.scalar_one_or_none()
    if cls is None:
        grade, _ = _infer_grade_level(class_name)
        cls = Class(name=class_name, grade_level=grade)
        db.add(cls)
        await db.flush()
        await db.refresh(cls)
    return cls


_TRUE_WORDS = {"是", "y", "yes", "true", "1", "班主任"}


def _parse_bool(raw: str) -> bool:
    return (raw or "").strip().lower() in _TRUE_WORDS


async def _homeroom_teacher_of_class(db: AsyncSession, cls: Class) -> User | None:
    """返回该班级的班主任：优先 Class.head_teacher_id，其次教师表中
    class_id + is_homeroom 的记录。"""
    if cls.head_teacher_id is not None:
        teacher = await db.get(User, cls.head_teacher_id)
        if teacher is not None and teacher.role == UserRole.TEACHER:
            return teacher
    result = await db.execute(
        select(User).where(
            User.role == UserRole.TEACHER,
            User.class_id == cls.id,
            User.is_homeroom.is_(True),
        )
    )
    return result.scalars().first()


async def _resolve_head_teacher(
    db: AsyncSession, identifier: str, class_id: int | None = None
) -> User | None:
    identifier = identifier.strip()
    if not identifier:
        return None
    # 手机号唯一匹配
    result = await db.execute(
        select(User).where(
            User.role == UserRole.TEACHER,
            User.phone == identifier,
        )
    )
    user = result.scalar_one_or_none()
    if user:
        return user
    # 同名教师常见，先在本班范围内按姓名匹配
    if class_id is not None:
        result = await db.execute(
            select(User).where(
                User.role == UserRole.TEACHER,
                User.real_name == identifier,
                User.class_id == class_id,
            )
        )
        candidates = result.scalars().all()
        if len(candidates) == 1:
            return candidates[0]
    # 再全局按姓名匹配——仅当唯一命中时返回，避免同名误配
    result = await db.execute(
        select(User).where(
            User.role == UserRole.TEACHER,
            User.real_name == identifier,
        )
    )
    candidates = result.scalars().all()
    return candidates[0] if len(candidates) == 1 else None



def _pick_sheet(wb, kind: str):
    """按导入类型挑选工作表：优先表名匹配 kind（中/英），否则退回活动表。"""
    want = {"teachers": ("teacher", "教师", "老师"), "students": ("student", "学生")}.get(kind, ())
    for ws in wb.worksheets:
        title = str(ws.title).lower()
        if any(k in title for k in want):
            return ws
    return wb.active


def _read_file(file: UploadFile, kind: str = "") -> list[dict[str, str]]:
    content = file.file.read()
    filename = (file.filename or "").lower()
    if filename.endswith(".csv"):
        text = content.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        return [dict(row) for row in reader]
    if filename.endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(content))
        ws = _pick_sheet(wb, kind)
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(c).strip() for c in rows[0]]
        return [
            {headers[i]: str(row[i]) if row[i] is not None else "" for i in range(len(headers))}
            for row in rows[1:]
        ]
    raise HTTPException(status_code=400, detail="仅支持 csv/xlsx")


async def _unique_username(db: AsyncSession, base: str) -> str:
    candidate = base
    suffix = 1
    while True:
        result = await db.execute(select(User).where(User.username == candidate))
        if result.scalar_one_or_none() is None:
            return candidate
        suffix += 1
        candidate = f"{base}{suffix}"


@router.get("/template/{kind}")
async def download_template(kind: str):
    if kind == "teachers":
        body = (
            "姓名,手机号,性别,班级,是否班主任\n"
            "张三,13800138000,男,七年级(1)班,是\n"
            "李四,13800138001,女,七年级(1)班,否\n"
        )
        filename = "teachers_template.csv"
    elif kind == "students":
        body = (
            "姓名,班级,性别,班主任\n"
            "王五,七年级(1)班,男,13800138000\n"
            "赵六,七年级(1)班,女,张三\n"
        )
        filename = "students_template.csv"
    else:
        raise HTTPException(status_code=400, detail="kind 必须是 teachers 或 students")

    from fastapi.responses import PlainTextResponse

    return PlainTextResponse(
        content=body,
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.post("/{kind}", response_model=ImportPreview)
async def preview_import(
    kind: str,
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    admin: User = Depends(require_admin),
):
    if kind not in ("teachers", "students"):
        raise HTTPException(status_code=400, detail="kind 必须是 teachers 或 students")
    rows = _read_file(file, kind)
    accounts = []
    errors = []
    warnings = []
    used_phones = set()

    for idx, row in enumerate(rows, start=2):
        name = row.get("姓名", "").strip()
        gender = row.get("性别", "").strip()
        if not name:
            errors.append(ImportRowError(row=idx, message="姓名不能为空"))
            continue
        if not gender:
            if kind == "teachers":
                gender = "未知"  # 教师名单常缺性别列
            else:
                errors.append(ImportRowError(row=idx, message="性别不能为空"))
                continue
        elif gender not in ("男", "女"):
            errors.append(ImportRowError(row=idx, message="性别必须是 男/女"))
            continue

        if kind == "teachers":
            phone = row.get("手机号", "").strip()
            if not phone:
                errors.append(ImportRowError(row=idx, message="教师手机号不能为空"))
                continue
            if phone in used_phones:
                errors.append(ImportRowError(row=idx, message=f"手机号 {phone} 重复"))
                continue
            used_phones.add(phone)
            class_name = _normalize_class_name(row.get("班级", ""))
            is_homeroom = _parse_bool(row.get("是否班主任", ""))
            cls = await _resolve_class(db, class_name) if class_name else None
            if is_homeroom and cls is None:
                errors.append(ImportRowError(row=idx, message="班主任必须填写所在班级"))
                continue
            base_user = f"t{_pinyin(name)}"
            accounts.append(
                {
                    "row": idx,
                    "real_name": name,
                    "gender": gender,
                    "phone": phone,
                    "class_id": cls.id if cls else None,
                    "class_name": cls.name if cls else "",
                    "is_homeroom": is_homeroom,
                    "role": "teacher",
                    "base_username": base_user,
                }
            )
        else:
            class_name = _normalize_class_name(row.get("班级", ""))
            head = row.get("班主任", "").strip()
            if not class_name:
                errors.append(ImportRowError(row=idx, message="班级不能为空"))
                continue
            grade, warn = _infer_grade_level(class_name)
            if warn:
                warnings.append(f"第 {idx} 行班级 '{class_name}' 无法推断学段，默认初中")
            cls = await _resolve_class(db, class_name)
            if head:
                head_teacher = await _resolve_head_teacher(db, head, class_id=cls.id)
                if head_teacher is None:
                    warnings.append(
                        f"第 {idx} 行班主任 '{head}' 未找到或同名不唯一，暂不设置班主任"
                    )
            else:
                # 未填写时沿用该班已登记的班主任（教师导入时的班级+是否班主任）
                head_teacher = await _homeroom_teacher_of_class(db, cls)
            base_user = f"s{cls.id}{_pinyin(name)}"
            accounts.append(
                {
                    "row": idx,
                    "real_name": name,
                    "gender": gender,
                    "class_id": cls.id,
                    "class_name": cls.name,
                    "head_teacher_id": head_teacher.id if head_teacher else None,
                    "head_teacher_name": head_teacher.real_name if head_teacher else "",
                    "is_homeroom": False,
                    "role": "student",
                    "base_username": base_user,
                }
            )

    return ImportPreview(
        kind=kind,
        total_rows=len(rows),
        accounts=accounts,
        warnings=warnings,
        errors=errors,
    )


@router.post("/{kind}/commit", response_model=ImportCommitResult)
async def commit_import(
    kind: str,
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    admin: User = Depends(require_admin),
):
    preview = await preview_import(kind, file, db, admin)
    if preview.errors:
        raise HTTPException(status_code=400, detail=preview.errors)

    created = 0
    rows_for_roster = []
    created_accounts: list[tuple[dict, User]] = []
    for acc in preview.accounts:
        username = await _unique_username(db, acc["base_username"])
        password = _gen_password()
        user = User(
            role=UserRole.TEACHER if kind == "teachers" else UserRole.STUDENT,
            username=username,
            password_hash=hash_password(password),
            initial_password=password,
            must_change_password=False,
            real_name=acc["real_name"],
            gender=acc["gender"],
            phone=acc.get("phone"),
            class_id=acc.get("class_id"),
            is_homeroom=bool(acc.get("is_homeroom", False)),
            is_active=True,
        )
        db.add(user)
        await db.flush()
        created_accounts.append((acc, user))
        rows_for_roster.append(
            {
                "username": username,
                "real_name": acc["real_name"],
                "password": password,
                "class_name": acc.get("class_name", ""),
                "is_homeroom": bool(acc.get("is_homeroom", False)),
            }
        )
        created += 1

    # 维护"班级 ↔ 班主任"对应关系：
    # 1) 本次导入的班主任教师接管其班级的 head_teacher_id（同名靠班级+是否班主任区分）
    # 2) 学生导入时，若班级尚无班主任，则按学生行提供的班主任补链
    for acc, user in created_accounts:
        if acc.get("class_id") is None:
            continue
        cls = await db.get(Class, acc["class_id"])
        if cls is None:
            continue
        if acc.get("is_homeroom"):
            cls.head_teacher_id = user.id
        elif acc.get("head_teacher_id") and cls.head_teacher_id is None:
            cls.head_teacher_id = acc["head_teacher_id"]

    await db.commit()

    if created:
        import_created_total.labels(kind).inc(created)

    roster_html = _render_roster_html(kind, rows_for_roster)
    return ImportCommitResult(
        kind=kind, created=created, errors=[], roster_html=roster_html
    )


def _render_roster_html(kind: str, rows: list[dict]) -> str:
    if kind == "students":
        header = "<tr><th>姓名</th><th>用户名</th><th>初始密码</th><th>班级</th></tr>"
    else:
        header = (
            "<tr><th>姓名</th><th>用户名</th><th>初始密码</th>"
            "<th>班级</th><th>是否班主任</th></tr>"
        )
    body = ""
    for r in rows:
        if kind == "students":
            extra = f"<td>{r.get('class_name', '')}</td>"
        else:
            mark = "是" if r.get("is_homeroom") else ""
            extra = f"<td>{r.get('class_name', '')}</td><td>{mark}</td>"
        body += (
            f"<tr><td>{r['real_name']}</td><td>{r['username']}</td>"
            f"<td>{r['password']}</td>{extra}</tr>"
        )
    return (
        "<!doctype html><html><head><meta charset='utf-8'>"
        "<style>@media print{body{font-size:12px;} table{width:100%; border-collapse:collapse;}} "
        "table,th,td{border:1px solid #ccc; padding:6px;}</style></head>"
        f"<body><h2>{'学生' if kind=='students' else '教师'}账号名单</h2>"
        f"<table>{header}{body}</table></body></html>"
    )
