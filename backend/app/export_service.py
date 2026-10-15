"""班级测评结果导出（问题 3）。

管理员与班主任共用同一套构建逻辑，区别只在「班级范围」：
- 管理员：任意班级，但仅含学生（绝不包含教师量表数据）。
- 班主任：仅限本人担任班主任的班级。

导出**包含未完成学生**（状态列标「未完成」，分数列留空），
这正是前端「还有 N 位同学未完成」提示的语义。
"""

from __future__ import annotations

import csv
import io
from datetime import datetime, timezone
from typing import Any
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models import Assessment, AssessmentStatus, Class, Result, ScaleDef, User, UserRole

# 只允许导出学生量表；教师量表（teacher_ref）永不进入导出。
EXPORT_SCALES = ("mht", "mssmhs")

_SEVERITY_LABEL = {"critical": "严重", "high": "高度关注", "info": "提示"}
_SEVERITY_RANK = {"critical": 3, "high": 2, "info": 1}
_XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_CSV_MEDIA = "text/csv; charset=utf-8"


class ExportError(ValueError):
    """参数或数据不合法（路由层转成 400/404）。"""


async def get_scale_payload(db: AsyncSession, scale_id: str) -> dict[str, Any]:
    if scale_id not in EXPORT_SCALES:
        raise ExportError("scale_id 必须是 mht 或 mssmhs")
    result = await db.execute(select(ScaleDef).where(ScaleDef.scale_id == scale_id))
    scale = result.scalar_one_or_none()
    if scale is None:
        raise ExportError("量表不存在")
    return scale.payload


def _order_key(assessment: Assessment) -> float:
    dt = assessment.completed_at or assessment.started_at
    if dt is None:
        return 0.0
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.timestamp()


def _fmt_dt(dt: datetime | None) -> str:
    if dt is None:
        return ""
    try:
        return dt.strftime("%Y-%m-%d %H:%M")
    except Exception:  # pragma: no cover - defensive
        return str(dt)


def _risk_label(flags: list[dict] | None) -> str:
    best: str | None = None
    for flag in flags or []:
        severity = flag.get("severity")
        if best is None or _SEVERITY_RANK.get(severity, 0) > _SEVERITY_RANK.get(best, 0):
            best = severity
    return _SEVERITY_LABEL.get(best, "无") if best else "无"


def _llm_conclusion(result: Result | None) -> str:
    """「LLM 测评结论」单元格：有结论写原文；模板回退追加「（模板生成）」；无结果留空。"""
    if result is None:
        return ""
    text = (result.llm_summary or "").strip()
    if result.fallback_used:
        text = f"{text}（模板生成）" if text else "（模板生成）"
    return text


def _mssmhs_overall_label(scores: dict) -> str:
    level = scores.get("overall_level")
    if level == "normal":
        return "正常"
    if level == "overall_issue":
        return "整体需关注"
    return level or ""


async def build_class_export(
    db: AsyncSession, cls: Class, scale_id: str
) -> tuple[list[str], list[list[Any]]]:
    """返回 (表头, 数据行)，包含该班全部学生（含未完成）。"""
    payload = await get_scale_payload(db, scale_id)
    scale_name = payload.get("name", scale_id)
    subscales = payload.get("subscales", [])

    students = (
        await db.execute(
            select(User)
            .where(User.class_id == cls.id, User.role == UserRole.STUDENT)
            .order_by(User.id)
        )
    ).scalars().all()
    student_ids = [s.id for s in students]

    latest: dict[int, Assessment] = {}
    attempt_no: dict[int, int] = {}
    if student_ids:
        assessments = (
            await db.execute(
                select(Assessment)
                .where(
                    Assessment.user_id.in_(student_ids),
                    Assessment.scale_id == scale_id,
                )
                .options(selectinload(Assessment.result))
            )
        ).scalars().all()
        grouped: dict[int, list[Assessment]] = {}
        for assessment in assessments:
            grouped.setdefault(assessment.user_id, []).append(assessment)
        for group in grouped.values():
            group.sort(key=_order_key)
            for idx, assessment in enumerate(group, start=1):
                attempt_no[assessment.id] = idx
            for assessment in group:
                if assessment.status != AssessmentStatus.COMPLETED:
                    continue
                current = latest.get(assessment.user_id)
                if current is None or _order_key(assessment) >= _order_key(current):
                    latest[assessment.user_id] = assessment

    # 列顺序（任务书问题 3）：姓名|班级|性别|状态|量表|第几次|有效性(效度)|总分/总均分|评级|预警级别|LLM 测评结论|完成时间
    # MSSMHS 在「总均分」之后保留 10 个因子列，再接「评级」。
    headers: list[str] = ["姓名", "班级", "性别", "状态", "量表", "第几次", "有效性(效度)"]
    if scale_id == "mht":
        headers += ["总分", "评级"]
    else:
        headers += ["总均分"] + [sub.get("name", sub.get("key")) for sub in subscales]
        headers += ["评级"]
    headers += ["预警级别", "LLM 测评结论", "完成时间"]

    rows: list[list[Any]] = []
    for student in students:
        assessment = latest.get(student.id)
        if assessment is None:
            # 未完成：身份信息照写，测评相关列（含 LLM 结论）留空。
            row: list[Any] = [student.real_name, cls.name, student.gender, "未完成", scale_name, ""]
            row += [""] * (len(headers) - len(row))
            rows.append(row)
            continue

        scores = (assessment.result.scores if assessment.result else {}) or {}
        is_valid = assessment.is_valid
        validity = "有效" if is_valid is not False else "无效"
        row = [
            student.real_name,
            cls.name,
            student.gender,
            "已完成",
            scale_name,
            attempt_no.get(assessment.id, ""),
        ]
        if scale_id == "mht":
            validity_score = scores.get("validity_score")
            row.append(
                f"{validity}（效度{validity_score}）" if validity_score is not None else validity
            )
            row.append(scores.get("total_score", ""))
            row.append(scores.get("total_label") or scores.get("band") or "")
        else:
            row.append(validity)
            row.append(scores.get("total_mean", ""))
            factors = scores.get("factors") or {}
            for sub in subscales:
                info = factors.get(sub.get("key")) or {}
                row.append(info.get("mean", ""))
            row.append(_mssmhs_overall_label(scores))
        flags = assessment.result.risk_flags if assessment.result else []
        row.append(_risk_label(flags))
        row.append(_llm_conclusion(assessment.result))
        completed_at = assessment.completed_at or (
            assessment.result.created_at if assessment.result else None
        )
        row.append(_fmt_dt(completed_at))
        rows.append(row)

    return headers, rows


def to_csv_bytes(headers: list[str], rows: list[list[Any]]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(headers)
    for row in rows:
        writer.writerow(row)
    return ("\ufeff" + buf.getvalue()).encode("utf-8")


def to_xlsx_bytes(headers: list[str], rows: list[list[Any]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "测评结果"
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    for index, header in enumerate(headers, start=1):
        widest = len(str(header))
        for row in rows:
            widest = max(widest, len(str(row[index - 1])))
        sheet.column_dimensions[get_column_letter(index)].width = min(max(10, widest + 2), 40)
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def render_export(
    cls: Class, scale_id: str, fmt: str, headers: list[str], rows: list[list[Any]]
) -> tuple[bytes, str, str]:
    """返回 (内容, media_type, Content-Disposition)。"""
    if fmt == "csv":
        content = to_csv_bytes(headers, rows)
        media = _CSV_MEDIA
    elif fmt == "xlsx":
        content = to_xlsx_bytes(headers, rows)
        media = _XLSX_MEDIA
    else:
        raise ExportError("format 必须是 xlsx 或 csv")
    filename = f"{cls.name}_{scale_id}_测评结果.{fmt}"
    disposition = f"attachment; filename*=UTF-8''{quote(filename)}"
    return content, media, disposition
